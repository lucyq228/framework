"""Reporting: chart-pack, build-excel-prototype, build-readout. Inputs are run artifacts only."""

from __future__ import annotations

import html
import json
import re
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from corr_core import charts
from corr_core.tools import tool
from corr_core.tools.common import ToolContext

HEADER_FILL = PatternFill("solid", fgColor="27251F")
TITLE_FONT = Font(bold=True, size=13)
HEADER_FONT = Font(bold=True, color="FFFFFF")


def _read(run_dir: Path, name: str, default=None):
    p = run_dir / name
    return json.loads(p.read_text(encoding="utf-8")) if p.exists() else default


def _results(run_dir: Path) -> dict[str, dict]:
    return {p.stem: json.loads(p.read_text(encoding="utf-8")) for p in sorted((run_dir / "results").glob("*.json"))}


def _sheet_name(name: str, used: set) -> str:
    base = re.sub(r"[\[\]\*\?/\\:]", "_", name)[:28]
    out, i = base, 1
    while out in used:
        i += 1
        out = f"{base[:26]}_{i}"
    used.add(out)
    return out


class SheetWriter:
    def __init__(self, ws):
        self.ws = ws
        self.row = 1

    def title(self, text: str):
        self.ws.cell(self.row, 1, text).font = TITLE_FONT
        self.row += 1

    def text(self, text: str, bold: bool = False):
        c = self.ws.cell(self.row, 1, text)
        c.font = Font(bold=bold)
        c.alignment = Alignment(wrap_text=False)
        self.row += 1

    def table(self, name: str | None, rows: list[dict]):
        if not rows:
            return
        if name:
            self.ws.cell(self.row, 1, name).font = Font(bold=True, italic=True)
            self.row += 1
        cols = list(dict.fromkeys(k for r in rows for k in r))
        for j, c in enumerate(cols, 1):
            cell = self.ws.cell(self.row, j, str(c))
            cell.font, cell.fill = HEADER_FONT, HEADER_FILL
        for r in rows:
            self.row += 1
            for j, c in enumerate(cols, 1):
                v = r.get(c)
                if isinstance(v, (list, dict)):
                    v = json.dumps(v, default=str)
                cell = self.ws.cell(self.row, j, v)
                if isinstance(v, float) and 0 <= abs(v) <= 1 and any(k in str(c) for k in ("rate", "share", "csat", "prog", "cons", "index", "kpi", "pass")):
                    cell.number_format = "0.0%"
                elif isinstance(v, float):
                    cell.number_format = "0.000"
        self.row += 2

    def image(self, path: str):
        p = Path(path)
        if p.exists():
            img = XLImage(str(p))
            img.width, img.height = img.width * .6, img.height * .6
            self.ws.add_image(img, f"A{self.row}")
            self.row += int(img.height / 20) + 2

    def autosize(self):
        for col in self.ws.columns:
            width = max((len(str(c.value)) for c in col[:200] if c.value is not None), default=8)
            self.ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(10, width + 2), 60)


@tool("chart-pack")
def chart_pack(tc: ToolContext, params: dict) -> dict:
    run = tc.run_dir
    found = sorted(str(p) for p in (run / "charts").glob("*.png"))
    sc = pd.read_parquet(run / "datasets" / "market_scorecard.parquet")
    sc = sc[sc["market"] != "All in scope"]
    labels = [f"{m} {k}" for m in sc["market"] for k in ("Taste", "Hosp")]
    vals = [v for _, r in sc.iterrows() for v in (r["taste_index_prog"], r["hosp_index_prog"])]
    found.append(charts.bar(run / "charts" / "scorecard__gse_indexes.png", labels, vals, "Gold Standard indexes by market (progressive)"))
    return {"summary": f"{len(found)} charts in the pack", "charts": found}


def _insights_or_fail(run: Path) -> dict:
    ins = _read(run, "insights.json")
    if not ins:
        raise FileNotFoundError("insights.json is missing; A7 must save insights before deliverables are built.")
    return ins


@tool("build-excel-prototype")
def build_excel_prototype(tc: ToolContext, params: dict) -> dict:
    run = tc.run_dir
    scope = tc.scope
    ins = _insights_or_fail(run)
    review = _read(run, "review.json", {})
    results = _results(run)
    wb = Workbook()
    used: set = set()

    ws = wb.active
    ws.title = _sheet_name("README", used)
    w = SheetWriter(ws)
    w.title(f"{scope['project_name']}: PACE-CSAT correlation prototype")
    w.text(ins.get("headline", ""), bold=True)
    w.text(f"Markets: {', '.join(scope['markets'])} | Period: {scope['period']['start']} to {scope['period']['end']} | "
           f"Primary scoring: {scope['primary_scoring']} | Primary grain: {scope['primary_grain']} | "
           f"CSAT window: {scope.get('primary_csat_window')}")
    w.text(f"Decisions: {json.dumps(scope['decisions'])}")
    w.text(f"Review verdict: {review.get('verdict', 'n/a')}")
    w.text("Synthetic data. Correlations are associations, not causal effects.")
    w.text("")
    w.text("Sheets: Insights, Scorecard, Descriptives, Hypotheses, one sheet per analysis, Restaurant Detail, Data Notes.")
    w.autosize()

    w = SheetWriter(wb.create_sheet(_sheet_name("Insights", used)))
    w.title("Insights")
    for s in ins.get("executive_summary", []):
        w.text(f"• {s}")
    w.row += 1
    w.table(None, [{"id": i["id"], "title": i["title"], "tag": i.get("tag"), "importance": i.get("importance"),
                    "robustness": i.get("robustness"), "so_what": i.get("so_what"),
                    "evidence": " | ".join(e["statement"] for e in i.get("evidence", []))} for i in ins.get("insights", [])])
    w.autosize()

    w = SheetWriter(wb.create_sheet(_sheet_name("Scorecard", used)))
    w.title("Gold Standard Execution scorecard (deck method)")
    sc = pd.read_parquet(run / "datasets" / "market_scorecard.parquet")
    w.table(None, sc.round(4).to_dict("records"))
    p = run / "charts" / "scorecard__gse_indexes.png"
    w.image(str(p))
    w.autosize()

    w = SheetWriter(wb.create_sheet(_sheet_name("Descriptives", used)))
    w.title("Descriptive EDA")
    eda = _read(run, "eda_descriptives.json", {})
    for step in eda.get("steps", []):
        w.text(f"{step['tool']}: {step['summary']}", bold=True)
        for o in step.get("observations", []):
            w.text(f"• {o}")
        for name, rows in (step.get("tables") or {}).items():
            w.table(name, rows)
        for c in step.get("charts", []):
            w.image(c)
    w.autosize()

    w = SheetWriter(wb.create_sheet(_sheet_name("Hypotheses", used)))
    w.title("Hypotheses and verdicts")
    verdicts = {h["id"]: h for h in review.get("hypothesis_verdicts", [])}
    w.table(None, [{"id": h["id"], "statement": h["statement"], "verdict": verdicts.get(h["id"], {}).get("verdict"),
                    "evidence": verdicts.get(h["id"], {}).get("reason"),
                    "tools": ", ".join(s["tool"] for s in h["pathway"])} for h in scope["hypotheses"]])
    w.autosize()

    for name, res in results.items():
        w = SheetWriter(wb.create_sheet(_sheet_name(name, used)))
        w.title(f"{res.get('tool')} ({name})")
        w.text(res.get("summary", ""), bold=True)
        if res.get("purpose"):
            w.text(f"Purpose: {res['purpose']}")
        w.text(f"Parameters: {json.dumps(res.get('params', {}), default=str)}")
        for n in res.get("notes", []):
            w.text(f"Note: {n}")
        w.row += 1
        if res.get("findings"):
            w.table("findings", res["findings"])
        for tname, rows in (res.get("tables") or {}).items():
            w.table(tname, rows)
        for c in res.get("charts", [])[:3]:
            w.image(c)
        w.autosize()

    w = SheetWriter(wb.create_sheet(_sheet_name("Restaurant Detail", used)))
    w.title("Restaurant-period detail (PACE, CSAT, indexes)")
    rp = pd.read_parquet(run / "datasets" / "restaurant_period.parquet")
    keep = [c for c in rp.columns if not c.startswith(("fail_",))]
    w.table(None, rp[keep].round(4).astype(object).where(rp[keep].notna(), None).to_dict("records"))
    w.autosize()

    w = SheetWriter(wb.create_sheet(_sheet_name("Data Notes", used)))
    w.title("Data quality, feasibility, and limitations")
    dq = _read(run, "dq_report.json", {})
    w.table("data quality checks", [{k: v for k, v in c.items() if k != "detail"} for c in dq.get("checks", [])])
    fz = _read(run, "feasibility_review.json", {})
    w.table("feasibility by pathway step", fz.get("steps", []))
    w.table("data requests", fz.get("data_requests", []))
    w.table("limitations", [{"limitation": x} for x in ins.get("limitations", [])])
    w.autosize()

    out = run / f"{re.sub(r'[^A-Za-z0-9]+', '_', scope['project_name']).strip('_')}_prototype.xlsx"
    wb.save(out)
    return {"summary": f"Excel prototype with {len(wb.sheetnames)} sheets", "file": out.name, "sheets": wb.sheetnames}


@tool("build-readout")
def build_readout(tc: ToolContext, params: dict) -> dict:
    run = tc.run_dir
    scope = tc.scope
    ins = _insights_or_fail(run)
    review = _read(run, "review.json", {})
    results = _results(run)
    lines = [f"# {scope['project_name']}: insight readout", "", f"**{ins.get('headline', '')}**", "",
             f"Scope: {', '.join(scope['markets'])}, {scope['period']['start']} to {scope['period']['end']}, "
             f"{scope['primary_scoring']} PACE scores, CSAT window {scope.get('primary_csat_window')}. "
             f"Review verdict: **{review.get('verdict', 'n/a')}**. Synthetic data.", "", "## Summary", ""]
    lines += [f"- {s}" for s in ins.get("executive_summary", [])]
    lines += ["", "## Insights", ""]
    for i in ins.get("insights", []):
        lines += [f"### {i['id']}. {i['title']}", "",
                  f"*{i.get('tag', '')} · importance {i.get('importance', '')} · robustness {i.get('robustness', '')}*", ""]
        for e in i.get("evidence", []):
            lines.append(f"- {e['statement']} (`{e['result']}`)")
        lines += ["", f"**So what:** {i.get('so_what', '')}", ""]
        for e in i.get("evidence", [])[:1]:
            for c in results.get(e["result"], {}).get("charts", [])[:1]:
                lines += [f"![{i['title']}](charts/{Path(c).name})", ""]
    for title, key in (("Limitations", "limitations"), ("Data requests", "data_requests"), ("Next steps", "next_steps")):
        items = ins.get(key, [])
        if items:
            lines += [f"## {title}", ""] + [f"- {x if isinstance(x, str) else json.dumps(x)}" for x in items] + [""]
    md = "\n".join(lines)
    (run / "insight_readout.md").write_text(md, encoding="utf-8")
    body = []
    for line in lines:
        if line.startswith("### "):
            body.append(f"<h3>{html.escape(line[4:])}</h3>")
        elif line.startswith("## "):
            body.append(f"<h2>{html.escape(line[3:])}</h2>")
        elif line.startswith("# "):
            body.append(f"<h1>{html.escape(line[2:])}</h1>")
        elif line.startswith("!["):
            src = line[line.index("(") + 1:-1]
            body.append(f'<img src="{src}" style="max-width:640px">')
        elif line.startswith("- "):
            body.append(f"<li>{_inline(line[2:])}</li>")
        elif line:
            body.append(f"<p>{_inline(line)}</p>")
    page = ("<html><head><meta charset='utf-8'><title>Insight readout</title><style>body{font-family:Segoe UI,Arial;"
            "max-width:860px;margin:2em auto;color:#27251F}h1{color:#DA291C}h3{margin-top:1.4em}</style></head><body>"
            + "\n".join(body) + "</body></html>")
    (run / "insight_readout.html").write_text(page, encoding="utf-8")
    return {"summary": f"Readout with {len(ins.get('insights', []))} insights", "files": ["insight_readout.md", "insight_readout.html"]}


def _inline(text: str) -> str:
    t = html.escape(text)
    t = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", t)
    t = re.sub(r"\*(.+?)\*", r"<i>\1</i>", t)
    return re.sub(r"`(.+?)`", r"<code>\1</code>", t)
